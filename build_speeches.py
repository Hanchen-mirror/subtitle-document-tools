import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import argparse
parser = argparse.ArgumentParser(description="将人工整理的 JSON 内容排成演讲稿 DOCX")
parser.add_argument("input", type=Path)
parser.add_argument("--output-dir", type=Path, default=Path("output"))
args = parser.parse_args()
OUT = args.output_dir.resolve()
OUT.mkdir(parents=True, exist_ok=True)
items = json.loads(args.input.read_text(encoding='utf-8-sig'))
for item in items:
    name = item['name']
    if not name or name in {'.', '..'} or any(c in name for c in '/\\:'):
        raise ValueError('name 必须是文件名，不可包含路径')
    if (OUT / (name + '.docx')).exists():
        raise FileExistsError('输出已存在，请指定其他目录或名称')

def font_style(style, east_asia, size, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for attr, value in [('ascii', 'Times New Roman'), ('hAnsi', 'Times New Roman'), ('eastAsia', east_asia)]:
        fonts.set(qn('w:' + attr), value)
    for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + attr), None)

for item in items:
    doc = Document()
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.2)
    sec.left_margin, sec.right_margin = Cm(2.6), Cm(2.6)
    sec.footer_distance = Cm(1.0)
    normal = doc.styles['Normal']
    font_style(normal, '宋体', 13)
    pf = normal.paragraph_format
    pf.line_spacing = 1.4
    pf.space_before, pf.space_after = Pt(0), Pt(5)
    pf.first_line_indent = Pt(26)
    pf.widow_control = True
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    font_style(doc.styles['Title'], '黑体', 20, True)
    tpf = doc.styles['Title'].paragraph_format
    tpf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tpf.first_line_indent = Pt(0)
    tpf.space_before, tpf.space_after = Pt(0), Pt(18)
    tpf.keep_with_next = True
    font_style(doc.styles['Heading 1'], '黑体', 16, True)
    hpf = doc.styles['Heading 1'].paragraph_format
    hpf.first_line_indent = Pt(0)
    hpf.space_before, hpf.space_after = Pt(0), Pt(12)
    hpf.keep_with_next = True
    doc.add_paragraph(item['title'], 'Title')
    for text in item['paragraphs']:
        doc.add_paragraph(text)
    if item.get('appendix'):
        heading = doc.add_paragraph(item['appendix_title'], 'Heading 1')
        heading.paragraph_format.page_break_before = True
        for text in item['appendix']:
            p = doc.add_paragraph(text)
            if text.startswith('评委'):
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.keep_with_next = True
                for run in p.runs:
                    run.bold = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    r = footer.add_run()
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0, 0, 0)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    field_r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = '1'
    field_r.append(t)
    field.append(field_r)
    footer._p.append(field)
    doc.core_properties.title = item['title']
    doc.core_properties.subject = '项目演讲稿'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.comments = ''
    path = OUT / (item['name'] + '.docx')
    doc.save(path)
    count = sum(len(t) for t in item['paragraphs'] + item.get('appendix', []))
    print(f'{path.name}: {count} characters, {len(doc.paragraphs)} paragraphs')
